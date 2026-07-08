"""Read a teacher's SAC gradebook spreadsheet into a structured model.

This is a *marks-already-exist* ingestion path (no scanning): a VCE teacher's
per-SAC workbook where each SAC lives on its own sheet as a question × student
grid. It powers `markable gradebook`, which renders a multi-SAC dashboard and,
with a study-design map, a per-student skills/content view.

Expected sheet shape (detected, not hard-coded to fixed rows):

    row h-1 :  ... | 1 | 2 | 3 | 10a | 10b | ... | TOTAL      (question labels)
    row h   :  ID | VCAA Number | Surname | First Name | ... | /2 | /4 | ... | /40
    row h+1+:  1  | 23173583X   | BLETSAS | Pauly       | ... |  2 |  2 | ...

The header row `h` is found by its "Surname" cell; question columns are the
cells on row `h` whose value looks like a max-mark (`/2`); the label sits one
row above. A SAC's number/topic/total come from the labelled cells near the top.

Requires the `xlsx` extra: ``uv sync --extra xlsx``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class GradeQuestion:
    id: str  # "1", "10a", or a criterion name like "Scientific Communication"
    max_marks: float


@dataclass
class GradeStudent:
    id: str
    surname: str
    first_name: str
    vcaa: str = ""
    class_group: str = ""
    marks: dict[str, float] = field(default_factory=dict)  # question id -> awarded

    @property
    def name(self) -> str:
        return f"{self.surname}, {self.first_name}".strip(", ")


@dataclass
class SAC:
    number: str
    topic: str
    total_marks: float
    sheet: str
    questions: list[GradeQuestion]
    students: list[GradeStudent]

    def total_for(self, s: GradeStudent) -> float:
        return sum(s.marks.get(q.id, 0.0) for q in self.questions)


@dataclass
class Gradebook:
    source: str
    sacs: list[SAC]
    mock_classes: bool = False  # True when class groups were assigned, not parsed


_MAXMARK = re.compile(r"^/\s*(\d+(?:\.\d+)?)$")


def _f(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _find_header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 40) + 1):
        for c in range(1, min(ws.max_column, 8) + 1):
            if str(ws.cell(r, c).value).strip().lower() == "surname":
                return r
    return None


def _labelled(ws, label: str, max_row: int = 8) -> str | None:
    """Value to the right of a `LABEL:` cell near the sheet top (SAC #, TOPIC…)."""
    target = label.lower().rstrip(":")
    for r in range(1, max_row + 1):
        for c in range(1, min(ws.max_column, 6) + 1):
            cell = str(ws.cell(r, c).value or "").strip().lower().rstrip(":")
            if cell == target:
                for cc in range(c + 1, c + 5):
                    v = ws.cell(r, cc).value
                    if v not in (None, ""):
                        return str(v).strip()
    return None


def _parse_sheet(ws) -> SAC | None:
    h = _find_header_row(ws)
    if h is None:
        return None

    # Which column holds Surname / First Name / ID / VCAA.
    cols = {}
    for c in range(1, ws.max_column + 1):
        key = str(ws.cell(h, c).value or "").strip().lower()
        if key in {"surname", "first name", "id", "vcaa number", "class", "form", "class group"}:
            cols[key] = c
    if "surname" not in cols:
        return None
    class_col = cols.get("class") or cols.get("class group") or cols.get("form")

    # Question columns: max-mark cells on the header row, label one row above.
    questions: list[GradeQuestion] = []
    q_cols: list[tuple[int, str]] = []
    for c in range(1, ws.max_column + 1):
        m = _MAXMARK.match(str(ws.cell(h, c).value or "").strip())
        if not m:
            continue
        label = ws.cell(h - 1, c).value
        label = str(label).strip() if label not in (None, "") else str(c)
        if label.upper() == "TOTAL":
            continue  # the TOTAL column is derived, not a question
        questions.append(GradeQuestion(id=label, max_marks=float(m.group(1))))
        q_cols.append((c, label))
    if not questions:
        return None

    students: list[GradeStudent] = []
    for r in range(h + 1, ws.max_row + 1):
        surname = ws.cell(r, cols["surname"]).value
        if surname in (None, ""):
            continue
        first = ws.cell(r, cols.get("first name", cols["surname"])).value
        sid = ws.cell(r, cols["id"]).value if "id" in cols else str(len(students) + 1)
        student = GradeStudent(
            id=str(sid).strip() if sid not in (None, "") else str(len(students) + 1),
            surname=str(surname).strip(),
            first_name=str(first or "").strip(),
            vcaa=str(ws.cell(r, cols["vcaa number"]).value or "").strip() if "vcaa number" in cols else "",
            class_group=str(ws.cell(r, class_col).value or "").strip() if class_col else "",
        )
        for c, label in q_cols:
            v = _f(ws.cell(r, c).value)
            if v is not None:
                student.marks[label] = v
        if student.marks:
            students.append(student)
    if not students:
        return None

    total = _f(_labelled(ws, "TOTAL MARKS")) or sum(q.max_marks for q in questions)
    return SAC(
        number=str(_labelled(ws, "SAC #") or ws.title),
        topic=_labelled(ws, "TOPIC") or ws.title,
        total_marks=total,
        sheet=ws.title,
        questions=questions,
        students=students,
    )


def read_gradebook(path: Path) -> Gradebook:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Reading gradebooks needs the 'xlsx' extra: uv sync --extra xlsx") from exc

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    sacs = []
    for ws in wb.worksheets:
        sac = _parse_sheet(ws)
        if sac is not None:
            sacs.append(sac)
    wb.close()
    if not sacs:
        raise ValueError(
            f"no SAC sheets found in {path.name} — expected sheets with a 'Surname' "
            "header row and '/N' max-mark columns"
        )
    return Gradebook(source=path.name, sacs=sacs)


def assign_mock_classes(gb: Gradebook, groups: list[str]) -> None:
    """Deterministically assign each student to one of `groups` when the workbook
    carries no class column — for demoing class-by-class views. Stable across
    SACs (keyed on the student's name)."""
    import hashlib

    gb.mock_classes = True
    for sac in gb.sacs:
        for st in sac.students:
            if st.class_group:
                continue
            h = int(hashlib.sha1(st.name.encode("utf-8")).hexdigest(), 16)
            st.class_group = groups[h % len(groups)]


def classes(gb: Gradebook) -> list[str]:
    seen = []
    for sac in gb.sacs:
        for st in sac.students:
            if st.class_group and st.class_group not in seen:
                seen.append(st.class_group)
    return sorted(seen)
